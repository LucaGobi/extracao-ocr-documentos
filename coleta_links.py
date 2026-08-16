import os
import time
import base64
import warnings
import re
import pyperclip # Biblioteca que lê o Ctrl+C
from datetime import datetime
import ctypes # <--- Adicionado para controle do Windows

# --- SILENCIADOR DE AVISOS ---
warnings.simplefilter(action='ignore', category=FutureWarning)
warnings.simplefilter(action='ignore', category=UserWarning)
os.environ['GRPC_VERBOSITY'] = 'ERROR'
os.environ['GLOG_minloglevel'] = '2'

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
import pytesseract
from PIL import Image
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, Inches

# ================= CONFIGURAÇÃO =================
# 1. COLE SUA API KEY AQUI:
GENAI_API_KEY = "" 
genai.configure(api_key=GENAI_API_KEY)

# 2. CAMINHO DO TESSERACT:
pytesseract.pytesseract.tesseract_cmd = r'' # <--- Insira o caminho local ao tesseract
# ================================================

# --- FUNÇÕES ANTI-SONO ---
def impedir_suspensao():
    """Impede o Windows de dormir enquanto o script roda."""
    print("SISTEMA: Windows nao ira suspender")
    ES_CONTINUOUS = 0x80000000
    ES_SYSTEM_REQUIRED = 0x00000001
    ES_DISPLAY_REQUIRED = 0x00000002
    ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS | ES_SYSTEM_REQUIRED | ES_DISPLAY_REQUIRED)

def permitir_suspensao():
    """Libera o Windows para dormir novamente."""
    print("SISTEMA: Windows liberado para suspensao")
    ES_CONTINUOUS = 0x80000000
    ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS)

def obter_caminho_seguro():
    """Cria/Retorna a pasta DOCUMENTOS_PROCESSADOS no local do script."""
    caminho_script = os.path.dirname(os.path.abspath(__file__))
    pasta_destino = os.path.join(caminho_script, "DOCUMENTOS_PROCESSADOS")
    if not os.path.exists(pasta_destino):
        os.makedirs(pasta_destino)
    return pasta_destino

def extrair_links_do_texto(texto_sujo):
    """Filtra links ..."""
    padrao = r"(\S+)" # <--- Insira o padrao de link a ser encontrado
    links_encontrados = re.findall(padrao, texto_sujo)
    return list(set(links_encontrados))

def formatar_markdown_para_word(paragrafo, texto):
    """Transforma **texto** em Negrito real no Word."""
    partes = texto.split('**')
    for i, parte in enumerate(partes):
        run = paragrafo.add_run(parte)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        if i % 2 != 0: 
            run.bold = True

def salvar_word_lote(lista_documentos):
    """Gera um único arquivo Word com todos os documentos."""
    if not lista_documentos: return None

    print(f"\nGerando arquivo Word com {len(lista_documentos)} documentos")
    doc = Document()
    
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    for i, texto_documento in enumerate(lista_documentos):
        p_header = doc.add_paragraph()
        run_h = p_header.add_run(f"DOCUMENTO PROCESSADO #{i+1}")
        run_h.bold = True
        run_h.font.size = Pt(8)
        
        linhas = texto_documento.split('\n')
        for linha in linhas:
            if linha.strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                formatar_markdown_para_word(p, linha)
        
        if i < len(lista_documentos) - 1:
            doc.add_page_break()

    nome_arquivo = f"LOTE_DOCUMENTOS_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    pasta_destino = obter_caminho_seguro()
    caminho_completo = os.path.join(pasta_destino, nome_arquivo)
    doc.save(caminho_completo)
    return caminho_completo

def encontrar_modelo_disponivel():
    try: return 'gemini-2.5-flash'
    except: return 'gemini-pro'

def formatar_com_ia(texto_bruto_ocr):
    nome_modelo = encontrar_modelo_disponivel()
    print(f"   -> Usando IA: {nome_modelo}")
    
    try:
        model = genai.GenerativeModel(nome_modelo)
        prompt_mestre = f"""
        Insira o prompt
        """
        response = model.generate_content(prompt_mestre)
        return response.text
    except Exception as e:
        return f"ERRO NA IA: {e}"

def processar_url_individual(driver, wait, url):
    texto_retorno = ""
    pasta_segura = obter_caminho_seguro()
    nome_arquivo_imagem = os.path.join(pasta_segura, "temp_work.png")
    
    try:
        print(f"\n--- Processando: {url} ---")
        driver.set_window_size(1920, 1080)
        driver.get(url)

        print("   -> Procurando botão")
        try:
            botao = wait.until(EC.presence_of_element_located((By.XPATH, "//*[contains(text(), 'Visualizar Documento')]")))
            driver.execute_script("arguments[0].click();", botao)
        except: return "ERRO: Botão não encontrado."

        num_abas = len(driver.window_handles)
        print("   -> Aguardando documento")
        nova_aba = False
        for _ in range(10):
            if len(driver.window_handles) > num_abas:
                driver.switch_to.window(driver.window_handles[-1])
                nova_aba = True
                break
            time.sleep(1)
        if not nova_aba: time.sleep(3)

        print("   -> Capturando imagem")
        time.sleep(3)
        try: h = driver.execute_script("return Math.max( document.body.scrollHeight, document.body.offsetHeight, document.documentElement.clientHeight, document.documentElement.scrollHeight, document.documentElement.offsetHeight );")
        except: h = 2500
        
        driver.set_window_size(1920, h + 300)
        time.sleep(1)
        driver.find_element(By.TAG_NAME, "body").screenshot(nome_arquivo_imagem)

        #Processa a imagem para facilitar a leitura
        img = Image.open(nome_arquivo_imagem).convert('L')
        img = img.resize((img.width * 2, img.height * 2), Image.Resampling.LANCZOS)
        img = img.point(lambda p: 0 if p < 160 else 255)
        
        print("   -> Lendo texto (OCR)")
        texto_ocr = pytesseract.image_to_string(img, lang='por')
        
        if len(texto_ocr) < 50: return "ERRO: Leitura de imagem falhou."
            
        print("   -> Formatando")
        texto_retorno = formatar_com_ia(texto_ocr)
        
        if len(driver.window_handles) > 1:
            driver.close()
            driver.switch_to.window(driver.window_handles[0])

    except Exception as e:
        return f"ERRO CRÍTICO NO LINK: {e}"
    finally:
        if os.path.exists(nome_arquivo_imagem):
            try: os.remove(nome_arquivo_imagem)
            except: pass
    return texto_retorno

def main():
    try:
        impedir_suspensao()

        print("\n=== LEITURA AUTOMÁTICA (CTRL+C) ===")
        print("1. Copie o texto/links da conversa (Ctrl + C).")
        print("2. Apenas aperte Enter aqui para processar o que está na memória.")
        print("="*60)
        input("Pressione ENTER para ler sua Área de Transferência")

        texto_colado = pyperclip.paste()
        
        print("\n--- TEXTO ENCONTRADO NA MEMÓRIA ---")
        print(texto_colado[:100] + "..." if len(texto_colado) > 100 else texto_colado)
        print("-----------------------------------")

        lista_urls = extrair_links_do_texto(texto_colado)
        
        if not lista_urls:
            print("\nNenhum link válido de documento encontrado.")
            print("Certifique-se de dar Ctrl+C no texto antes de rodar.")
            input("Enter para sair.")
            return

        print(f"\nENCONTRADOS {len(lista_urls)} LINKS DE DOCUMENTOS! INICIANDO...")
        
        print("\nIniciando navegador (Invisível)...")
        options = webdriver.ChromeOptions()
        options.add_argument("--disable-popup-blocking")
        options.add_argument("--headless=new")
        options.add_argument("--disable-dev-shm-usage") # Estabilidade
        options.add_argument("--no-sandbox") # Estabilidade
        
        driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
        wait = WebDriverWait(driver, 15)

        resultados = []

        try:
            for i, url in enumerate(lista_urls):
                print(f"\n[{i+1}/{len(lista_urls)}] Processando")
                resultado = processar_url_individual(driver, wait, url)
                resultados.append(resultado)
                print("   -> Concluído.")
            
            if resultados:
                caminho_word = salvar_word_lote(resultados)
                print("\n" + "="*40)
                print(f"SUCESSO TOTAL! ARQUIVO WORD SALVO EM:\n{caminho_word}")
                print("="*40 + "\n")
                try: os.startfile(caminho_word)
                except: pass
                
        finally:
            driver.quit()
    
    finally:
        permitir_suspensao()

if __name__ == "__main__":
    main()
